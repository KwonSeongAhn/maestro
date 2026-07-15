# -*- coding: utf-8 -*-
"""① 시스템 전체 각파이프 통합본  ② 발주 패키지 표지  ③ RFQ 견적요청서  (모두 한/중)."""
import os, math
from collections import OrderedDict
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DWG="/tmp/claude-0/-home-user-maestro/89feb1d7-5414-58cf-b499-d60f30aed29c/scratchpad/dwg"
OUTD="/tmp/claude-0/-home-user-maestro/89feb1d7-5414-58cf-b499-d60f30aed29c/scratchpad/out"
DATE="2026-07-14"

# ---------- DOCX helper factory ----------
class Doc:
    def __init__(self, lang):
        self.lang=lang; self.CJK="Malgun Gothic" if lang=="kr" else "Microsoft YaHei"
        d=Document(); self.d=d; s=d.sections[0]
        s.page_width=Mm(210); s.page_height=Mm(297)
        for m in ("top_margin","bottom_margin","left_margin","right_margin"): setattr(s,m,Mm(16))
        nf=d.styles["Normal"].font; nf.name="Calibri"; nf.size=Pt(10)
        d.styles["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"),self.CJK)
    def T(self,kr,cn): return kr if self.lang=="kr" else cn
    def _cjk(self,run,size=10,bold=False,color=None):
        run.font.name="Calibri"; run.font.size=Pt(size); run.font.bold=bold
        if color: run.font.color.rgb=RGBColor(*color)
        rpr=run._element.get_or_add_rPr(); rf=rpr.find(qn("w:rFonts"))
        if rf is None: rf=OxmlElement("w:rFonts"); rpr.append(rf)
        rf.set(qn("w:eastAsia"),self.CJK)
    def para(self,t,s=10,b=False,c=None,after=4,al=None):
        p=self.d.add_paragraph(); p.paragraph_format.space_after=Pt(after)
        if al: p.alignment=al
        self._cjk(p.add_run(t),s,b,c); return p
    def h1(self,n,t):
        p=self.d.add_paragraph(); p.paragraph_format.space_before=Pt(9); p.paragraph_format.space_after=Pt(4)
        self._cjk(p.add_run(f"{n}. {t}"),13.5,True,(0x11,0x33,0x55))
        pr=p._p.get_or_add_pPr(); pb=OxmlElement("w:pBdr"); bt=OxmlElement("w:bottom")
        for k,v in (("w:val","single"),("w:sz","12"),("w:space","2"),("w:color","2B6CB0")): bt.set(qn(k),v)
        pb.append(bt); pr.append(pb)
    def title(self,t,sub=None):
        self.para(t,20,True,(0x11,0x33,0x55),2)
        if sub: self.para(sub,12,False,(0x55,0x55,0x55),8)
    def table(self,heads,rows,ctr=set(),fs=8.5,hdr=True):
        t=self.d.add_table(rows=1 if hdr else 0,cols=len(heads)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
        if hdr:
            hc=t.rows[0].cells
            for i,h in enumerate(heads):
                hc[i].text=""; self._cjk(hc[i].paragraphs[0].add_run(h),fs,True,(0xff,0xff,0xff))
                hc[i].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
                sh=OxmlElement("w:shd"); sh.set(qn("w:fill"),"1F2937"); hc[i]._tc.get_or_add_tcPr().append(sh)
        for row in rows:
            cells=t.add_row().cells
            for i,v in enumerate(row):
                cells[i].text=""; self._cjk(cells[i].paragraphs[0].add_run(str(v)),fs)
                cells[i].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER if i in ctr else WD_ALIGN_PARAGRAPH.LEFT
        return t
    def img(self,path,w=168,cap=None):
        self.d.add_picture(path,width=Mm(w)); self.d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
        if cap: self.para(cap,8.5,False,(0x66,0x66,0x66),6,WD_ALIGN_PARAGRAPH.CENTER)
    def save(self,name):
        p=f"{OUTD}/{name}.docx"; self.d.save(p); print("DOCX:",p)

def xstyle():
    thin=Side(style="thin",color="9AA4B2")
    return dict(bd=Border(left=thin,right=thin,top=thin,bottom=thin),
                hf=PatternFill("solid",fgColor="1F2937"), gf=PatternFill("solid",fgColor="E8EEF5"),
                tf=PatternFill("solid",fgColor="DFF0DF"),
                ctr=Alignment(horizontal="center",vertical="center",wrap_text=True),
                lft=Alignment(horizontal="left",vertical="center",wrap_text=True))

# =====================================================================
#  DATA — 시스템 전 각파이프 (route, sub_kr, sub_cn, partno, kr, cn, sec, mat, L, qty, stock)
# =====================================================================
A="호퍼 제작업체 (턴키)"; A2="料斗制作商(交钥匙)"
B="검사대/작업대 (별도 벤더)"; B2="检查台/工作台(另供)"
C="차광박스 (자가제작)"; C2="遮光箱(自制)"
SYS=[
 # A. 호퍼
 (A,A2,"LP-01","베이스 전/후빔","底架前/后梁","□60×60×2.3t","STK400",1528,2,6000),
 (A,A2,"LP-02","베이스 사이드빔","底架侧梁","□60×60×2.3t","STK400",1692,2,6000),
 (A,A2,"LP-03","호퍼받침 가로바","承托横杆","□60×60×2.3t","STK400",1528,1,6000),
 (A,A2,"LP-04","상단링 전/후빔","上环前/后梁","□60×60×2.3t","STK400",1528,2,6000),
 (A,A2,"LP-05","상단링 사이드빔","上环侧梁","□60×60×2.3t","STK400",1692,2,6000),
 (A,A2,"LP-06","수직 다리","立柱腿","□60×60×2.3t","STK400",257,7,6000),
 (A,A2,"LP-07","하부보강 기둥","下部加强柱","□50×50×2.3t","STK400",610,2,6000),
 (A,A2,"LP-08","상부보강 각봉","上部加强方钢","■50×50 SOLID","STK400",242,2,1000),
 (A,A2,"LP-09","결합부 직사각관","接合矩形管","□30×60×2.3t","STK400",1252,2,6000),
 (A,A2,"LP-10","코너 포스트","角立柱","□60×60×2.3t","SUS201",792,4,6000),
 (A,A2,"LP-11","림 전면","上缘前","□60×60×2.3t","SUS201",638,2,6000),
 (A,A2,"LP-12","림 후면","上缘后","□60×60×2.3t","SUS201",1464,1,6000),
 (A,A2,"LP-13","림 사이드","上缘侧","□60×60×2.3t","SUS201",1692,2,6000),
 # B. 검사대/작업대
 (B,B2,"WT-01","작업대 다리","工作台腿","□40×40×2.0t","STK400",955,4,6000),
 (B,B2,"WT-02","상부 가로대(장)","上横杆(长)","□40×60×2.0t","STK400",1290,2,6000),
 (B,B2,"WT-03","상부 가로대(단)","上横杆(短)","□40×60×2.0t","STK400",1040,2,6000),
 (B,B2,"WT-04","하부 보강대(장)","下加强杆(长)","□40×40×2.0t","STK400",1290,2,6000),
 (B,B2,"WT-05","하부 보강대(단)","下加强杆(短)","□40×40×2.0t","STK400",1040,2,6000),
 # C. 차광박스
 (C,C2,"SB-01","프레임 수평-장","框架水平-长","□30×30×1.4t","STEEL",980,4,6000),
 (C,C2,"SB-02","프레임 수평-단","框架水平-短","□30×30×1.4t","STEEL",540,4,6000),
 (C,C2,"SB-03","프레임 수직","框架竖","□30×30×1.4t","STEEL",896,4,6000),
]
def grand(items): return round(sum(r[7]*r[8] for r in items)/1000,2)
# 조달요약 by (sec,mat)
def summary(items):
    g=OrderedDict()
    for r in items:
        k=(r[5],r[6],r[9]); g.setdefault(k,0); g[k]+=r[7]*r[8]
    out=[]
    for (sec,mat,stock),tot in g.items():
        bars=max(1,math.ceil(tot*1.10/stock))
        out.append((sec,mat,round(tot/1000,2),stock/1000,bars,round(bars*stock/1000,1)))
    return out

# =====================================================================
# ① 시스템 전체 각파이프 통합본
# =====================================================================
GRAND_SYS=grand(SYS); SUM=summary(SYS)
routes=OrderedDict()
for r in SYS: routes.setdefault((r[0],r[1]),[]).append(r)

def build_sys(lang):
    D=Doc(lang); T=D.T
    D.title(T("시스템 각파이프 종합 조달 명세 (통합본)","系统方管综合采购清单(汇总)"),
            T("전 서브시스템 금속 각파이프 · 조달경로별  (LBH-PIPE-SYS-001)","全子系统金属方管·按采购路径 (LBH-PIPE-SYS-001)"))
    D.table([T("항목","项目"),T("내용","内容")],
      [[T("문서번호","文件号"),"LBH-PIPE-SYS-001"],[T("일자","日期"),DATE],
       [T("범위","范围"),T("호퍼+검사대/작업대+차광박스 금속 각파이프 전량","料斗+检查台/工作台+遮光箱 金属方管全量")],
       [T("총 연장","总延长"),f"{GRAND_SYS} m"],
       [T("제외","不含"),T("카메라 리그 □24(3D프린팅), 매달림봉 Ø8(환봉)","相机支架□24(3D打印),吊杆Ø8(圆棒)")]],ctr={0})
    D.h1(1,T("조달경로별 요약","按采购路径汇总"))
    rows=[]
    for (rk,rc),items in routes.items():
        rows.append([T(rk,rc),f"{len([1 for _ in items])} "+T("품목","项")+f" / {sum(i[8] for i in items)} "+T("본","根"),f"{grand(items):.2f} m"])
    D.table([T("조달경로","采购路径"),T("품목/수량","项/数量"),T("연장","延长")],rows,ctr={1,2})
    D.h1(2,T("재질·규격","材质·规格"))
    D.table([T("구분","分类"),T("사양","规格")],
     [[T("호퍼(A)","料斗(A)"),"STK400 □60/□50/□30×60, ■50 SOLID + SUS201 □60×60×2.3t"],
      [T("검사대(B)","检查台(B)"),T("STK400 □40×40×2.0t, □40×60×2.0t","STK400 □40×40×2.0t, □40×60×2.0t")],
      [T("차광박스(C)","遮光箱(C)"),T("스틸 각관 □30×30×1.4t, 회색 분체도장","钢方管□30×30×1.4t,灰色粉末涂装")],
      [T("마감","表面"),T("STK400·스틸=쇼트+분체도장 / SUS201=무처리","STK400·钢=喷砂+粉末涂装 / SUS201=不处理")]],ctr={0})
    D.para(T("※ 차광박스 프레임 재질은 최신 설계 기준 스틸 각관이나, 이력상 알루미늄 프로파일(§04)·라왕각재(초기 BOM)와 상이 — 자가제작 부재이므로 현장 확정.",
             "※ 遮光箱框架材质按最新设计为钢方管,但历史上与铝型材(§04)·木方(初期BOM)不一致 — 属自制件,现场确定。"),8.5,False,(0xc0,0x39,0x2b))
    D.d.add_page_break()
    D.h1(3,T("통합 절단표 (조달경로별)","汇总下料表(按路径)"))
    for (rk,rc),items in routes.items():
        D.para(T(rk,rc),11,True,(0x2b,0x6c,0xb0),3)
        rows=[[i[2],T(i[3],i[4]),i[5],i[6],i[7],i[8]] for i in items]
        rows.append([T("소계","小计"),"","","",f"{grand(items)}m",sum(i[8] for i in items)])
        D.table([T("품번","品番"),T("품명","品名"),T("단면","截面"),T("재질","材质"),T("길이","长度"),T("수량","数量")],rows,ctr={0,2,3,4,5},fs=8)
    D.h1(4,T("정척 조달 요약","定尺采购汇总"))
    rows=[[s,m,f"{tot:.2f}",f"{st:.0f}",b,f"{o:.1f}"] for s,m,tot,st,b,o in SUM]
    rows.append([T("합계","合计"),"",f"{GRAND_SYS:.2f}","","",""])
    D.table([T("단면","截面"),T("재질","材质"),T("총소요(m)","需求(m)"),T("정척(m)","定尺(m)"),T("본수","根数"),T("발주(m)","订购(m)")],rows,ctr={0,1,2,3,4,5},fs=9)
    D.save(f"LBH-PIPE-SYS-001_{'KR' if lang=='kr' else 'CN'}")

def xlsx_sys():
    S=xstyle(); wb=openpyxl.Workbook(); ws=wb.active; ws.title="통합절단표"
    ws.merge_cells("A1:H1"); ws["A1"]="시스템 각파이프 통합 절단표 / 系统方管汇总下料表 (LBH-PIPE-SYS-001)"; ws["A1"].font=Font(bold=True,size=12); ws["A1"].alignment=S["ctr"]
    heads=["조달경로 路径","품번","품명(한)","品名(中)","단면 截面","재질 材质","길이 长度","수량 数量"]
    for c,h in enumerate(heads,1):
        cell=ws.cell(3,c,h); cell.font=Font(bold=True,color="FFFFFF"); cell.fill=S["hf"]; cell.alignment=S["ctr"]; cell.border=S["bd"]
    r=4
    for row in SYS:
        vals=[f"{row[0]}\n{row[1]}",row[2],row[3],row[4],row[5],row[6],row[7],row[8]]
        for c,v in enumerate(vals,1):
            cell=ws.cell(r,c,v); cell.border=S["bd"]; cell.font=Font(size=9); cell.alignment=S["ctr"] if c in(2,5,6,7,8) else S["lft"]
        r+=1
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6); ws.cell(r,1,"총 연장 总延长 (m)").font=Font(bold=True); ws.cell(r,1).fill=S["tf"]; ws.cell(r,1).alignment=S["ctr"]
    ws.cell(r,7,GRAND_SYS).font=Font(bold=True); ws.cell(r,7).fill=S["tf"]; ws.cell(r,7).alignment=S["ctr"]; ws.cell(r,8,"").fill=S["tf"]
    for c in range(1,9): ws.cell(r,c).border=S["bd"]
    for i,w in enumerate([20,10,18,16,15,10,9,8],1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
    ws.row_dimensions[3].height=26; ws.freeze_panes="A4"
    # summary sheet
    ws2=wb.create_sheet("조달요약")
    ws2.merge_cells("A1:F1"); ws2["A1"]="정척 조달 요약 / 定尺采购汇总"; ws2["A1"].font=Font(bold=True,size=12); ws2["A1"].alignment=S["ctr"]
    for c,h in enumerate(["단면 截面","재질 材质","총소요(m)","정척(m)","본수 根数","발주(m)"],1):
        cell=ws2.cell(3,c,h); cell.font=Font(bold=True,color="FFFFFF"); cell.fill=S["hf"]; cell.alignment=S["ctr"]; cell.border=S["bd"]
    r=4
    for s,m,tot,st,b,o in SUM:
        for c,v in enumerate([s,m,tot,st,b,o],1):
            cell=ws2.cell(r,c,v); cell.border=S["bd"]; cell.alignment=S["ctr"]; cell.font=Font(size=10)
        r+=1
    for i,w in enumerate([16,12,12,10,10,12],1): ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
    p=f"{OUTD}/LBH-PIPE-SYS-001_CUTLIST.xlsx"; wb.save(p); print("XLSX:",p)

# =====================================================================
# ② 발주 패키지 표지 (커버시트)
# =====================================================================
def build_cover(lang):
    D=Doc(lang); T=D.T
    D.title(T("발주 패키지 표지","订购包封面"),T("호퍼 제작 일괄 발주 — 수신: 호퍼 제작업체","料斗制作一并订购 — 收件:料斗制作商"))
    D.table([T("구분","分类"),T("내용","内容")],
      [[T("프로젝트","项目"),T("LOSTBALL 로스트볼 선별설비 — 호퍼 어셈블리","LOSTBALL 高尔夫球分选设备 — 料斗总成")],
       [T("발주자","甲方"),T("(기입)","(填写)")],[T("수신(제작업체)","收件(制作商)"),T("(기입)","(填写)")],
       [T("발주일","订购日"),DATE],[T("납품지","交货地"),T("(기입)","(填写)")],
       [T("보안 구분","保密"),T("엘레베이터와 분리 발주 — 본 패키지는 호퍼 단독 완결","与提升机分开订购 — 本包为料斗独立完整")]],ctr={0})
    D.h1(1,T("동봉 문서","随附文件"))
    D.table([T("No","No"),T("문서번호","文件号"),T("제목","标题"),T("Rev","Rev"),T("구성","构成")],
      [["1","LBH-FAB-001",T("호퍼 제작 사양서(철판·캐스터·보강)","料斗制作规格书(钢板·脚轮·加强)"),"A",T("docx 한/중 + BOM xlsx","docx中韩+BOM xlsx")],
       ["2","LBH-PIPE-001",T("호퍼 각파이프 발주서","料斗方管订购书"),"A",T("docx 한/중 + 절단표 xlsx","docx中韩+下料表 xlsx")]],ctr={0,1,3})
    D.h1(2,T("발주 요지 (턴키 범위)","订购要旨(交钥匙范围)"))
    D.para(T("호퍼 제작업체가 아래 일체를 사입·가공·조립하여 완제 호퍼 1대를 납품한다:\n"
             "① 철판 SUS201 사입 → 전개도(LBH-FAB-001 §3) 절단·절곡·용접\n"
             "② 각파이프 STK400/SUS201 사입 → 절단표(LBH-PIPE-001) 절단·용접\n"
             "③ STK400 부재 쇼트+분체도장(RAL7016)\n"
             "④ 캐스터 Ø125 200kg ×8 볼트온 장착(마운트판 용접)\n"
             "⑤ 결합부 보강·고정절편·볼트류 조립 → 시운전(구름·제동) 후 포장·출하",
             "料斗制作商自购·加工·组装,交付成品料斗1台:\n"
             "① 钢板SUS201自购→按展开图(LBH-FAB-001 §3)下料·折弯·焊接\n"
             "② 方管STK400/SUS201自购→按下料表(LBH-PIPE-001)下料·焊接\n"
             "③ STK400件喷砂+粉末涂装(RAL7016)\n"
             "④ 脚轮Ø125 200kg×8 螺栓安装(安装板焊接)\n"
             "⑤ 接合加强·固定片·螺栓组装→试运行(滚动·制动)后包装·出货"),9.5)
    D.h1(3,T("확인·서명","确认·签署"))
    D.table([T("발주자 甲方","甲方"),T("제작업체 制作商","制作商")],[["\n\n(서명/날인)","\n\n(签字/盖章)"]],fs=10)
    D.save(f"LBH-COVER-001_{'KR' if lang=='kr' else 'CN'}")

# =====================================================================
# ③ RFQ 견적요청서
# =====================================================================
RFQ_ITEMS=[
 ("1",("호퍼 철판 가공·용접 일체 (SUS201 t4.0, 전개도 12판)","料斗钢板加工焊接(SUS201 t4.0,展开12板)"),"1 set"),
 ("2",("각파이프 사입·절단·용접 (STK400/SUS201, 총 29.71m)","方管自购下料焊接(STK400/SUS201,共29.71m)"),"1 set"),
 ("3",("보강재·고정절편·결합부 (STK400 t6~8, 50/30×60각)","加强件·固定片·接合(STK400 t6~8)"),"1 set"),
 ("4",("표면처리 쇼트+분체도장 RAL7016 (STK400 부재)","表面喷砂+粉末涂装RAL7016(STK400件)"),"1 set"),
 ("5",("캐스터 Ø125 200kg 선회제동 ×8 (지입 or 사입)","脚轮Ø125 200kg×8"),"8 EA"),
 ("6",("체결부품 A2-70 SHCS/너트/와셔 일체","紧固件A2-70 一套"),"1 set"),
 ("7",("조립·시운전·포장·내륙운송","组装·试运行·包装·内陆运输"),"1 set"),
]
def build_rfq(lang):
    D=Doc(lang); T=D.T
    D.title(T("견적요청서 (RFQ)","询价书(RFQ)"),T("호퍼 어셈블리 1대 — LBH-FAB-001 / LBH-PIPE-001 기준","料斗总成1台 — 依据LBH-FAB-001/LBH-PIPE-001"))
    D.table([T("항목","项目"),T("내용","内容")],
     [[T("RFQ 번호","询价号"),"LBH-RFQ-001"],[T("발행일","发行日"),DATE],
      [T("발주자","甲方"),T("(기입)","(填写)")],[T("수신","收件"),T("(제작업체명)","(制作商名)")],
      [T("회신 기한","回复期限"),T("(기입)","(填写)")],[T("문의","咨询"),"kremlin4277@gmail.com"]],ctr={0})
    D.h1(1,T("견적 대상 (수량 확정, 단가·납기 회신)","询价对象(数量确定,回复单价·交期)"))
    rows=[[n,T(kr,cn),u,"",""] for n,(kr,cn),u in RFQ_ITEMS]
    rows.append([T("합계","合计"),"","","",""])
    D.table([T("No","No"),T("품목·사양","项目·规格"),T("수량","数量"),T("단가(USD) 单价","单价"),T("금액(USD) 金额","金额")],rows,ctr={0,2,3,4},fs=8.5)
    D.h1(2,T("견적 조건 (회신 기입)","报价条件(请填写)"))
    D.table([T("조건","条件"),T("회신","回复")],
     [[T("통화","货币"),"USD / CNY (   )"],[T("인코텀즈","贸易术语"),"FOB / CIF (   )  ______"],
      [T("납기(발주 후)","交期(下单后)"),"______ "+T("일","天")],[T("결제조건","付款"),"T/T  ___% "+T("선금","预付")+" / ___% "+T("잔금","尾款")],
      [T("견적 유효기간","报价有效期"),"______ "+T("일","天")],[T("보증","质保"),"______ "+T("개월","个月")],
      [T("검사","检验"),T("치수·용접·도장 검사성적서 제출","尺寸·焊接·涂装检验报告")],
      [T("포장","包装"),T("목재 크레이트, 방청·단면보호","木箱,防锈·端面保护")]],ctr={0})
    D.h1(3,T("첨부·주기","附件·备注"))
    D.para(T("· 상세 사양·도면·BOM: LBH-FAB-001(철판) + LBH-PIPE-001(각파이프) 참조.\n"
             "· 각파이프·철판 원자재는 제작업체 사입 기준. 캐스터는 지입/사입 각각 견적 병기 가능.\n"
             "· 엘레베이터는 별도 발주(본 RFQ 미포함).",
             "· 详细规格·图纸·BOM: 见LBH-FAB-001(钢板)+LBH-PIPE-001(方管)。\n"
             "· 方管·钢板原材由制作商自购。脚轮可甲供/自购分别报价。\n"
             "· 提升机另行订购(不含本RFQ)。"),9.5)
    D.save(f"LBH-RFQ-001_{'KR' if lang=='kr' else 'CN'}")

def xlsx_rfq():
    S=xstyle(); wb=openpyxl.Workbook(); ws=wb.active; ws.title="견적서 报价"
    ws.merge_cells("A1:F1"); ws["A1"]="견적요청 RFQ / 询价  LBH-RFQ-001  (단가·금액·납기 회신 记入)"; ws["A1"].font=Font(bold=True,size=12); ws["A1"].alignment=S["ctr"]
    for c,h in enumerate(["No","품목·사양 项目·规格","수량 数量","단가 单价","금액 金额","납기·비고 交期·备注"],1):
        cell=ws.cell(3,c,h); cell.font=Font(bold=True,color="FFFFFF"); cell.fill=S["hf"]; cell.alignment=S["ctr"]; cell.border=S["bd"]
    r=4
    for n,(kr,cn),u in RFQ_ITEMS:
        for c,v in enumerate([n,f"{kr}\n{cn}",u,"","",""],1):
            cell=ws.cell(r,c,v); cell.border=S["bd"]; cell.font=Font(size=9); cell.alignment=S["ctr"] if c in(1,3) else S["lft"]
        r+=1
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=4); ws.cell(r,1,"합계 合计 (USD)").font=Font(bold=True); ws.cell(r,1).fill=S["tf"]; ws.cell(r,1).alignment=S["ctr"]
    for c in (5,6): ws.cell(r,c,"").fill=S["tf"]
    for c in range(1,7): ws.cell(r,c).border=S["bd"]
    for i,w in enumerate([4,40,10,14,14,20],1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
    ws.row_dimensions[3].height=26
    p=f"{OUTD}/LBH-RFQ-001_QUOTE.xlsx"; wb.save(p); print("XLSX:",p)

# ---- run all ----
build_sys("kr"); build_sys("cn"); xlsx_sys()
build_cover("kr"); build_cover("cn")
build_rfq("kr"); build_rfq("cn"); xlsx_rfq()
print("SYS grand:",GRAND_SYS,"m | routes:",[(k[0],grand(v)) for k,v in routes.items()])
