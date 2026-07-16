# -*- coding: utf-8 -*-
"""LBH-3DP-001 Rev.B — 최종 62부품(홀·소켓·재구성) 3D프린트 문서 (한/중)."""
import json,os
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
SP="/tmp/claude-0/-home-user-maestro/89feb1d7-5414-58cf-b499-d60f30aed29c/scratchpad"; OUTD=f"{SP}/out"
MAN=json.load(open(f"{SP}/parts_final3/parts_manifest.json"))
tot_q=sum(x["qty"] for x in MAN)
bymat={}
for x in MAN:
    m=x["material_hint"]; bymat.setdefault(m,{"n":0,"q":0}); bymat[m]["n"]+=1; bymat[m]["q"]+=x["qty"]

class Doc:
    def __init__(s,lang):
        s.lang=lang; s.CJK="Malgun Gothic" if lang=="kr" else "Microsoft YaHei"
        d=Document(); s.d=d; sec=d.sections[0]; sec.page_width=Mm(210); sec.page_height=Mm(297)
        for m in("top_margin","bottom_margin","left_margin","right_margin"): setattr(sec,m,Mm(15))
        nf=d.styles["Normal"].font; nf.name="Calibri"; nf.size=Pt(10)
        d.styles["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"),s.CJK)
    def T(s,kr,cn): return kr if s.lang=="kr" else cn
    def _c(s,r,sz=10,b=False,col=None):
        r.font.name="Calibri"; r.font.size=Pt(sz); r.font.bold=b
        if col: r.font.color.rgb=RGBColor(*col)
        rp=r._element.get_or_add_rPr(); rf=rp.find(qn("w:rFonts"))
        if rf is None: rf=OxmlElement("w:rFonts"); rp.append(rf)
        rf.set(qn("w:eastAsia"),s.CJK)
    def para(s,t,sz=10,b=False,col=None,after=4):
        p=s.d.add_paragraph(); p.paragraph_format.space_after=Pt(after); s._c(p.add_run(t),sz,b,col); return p
    def h1(s,n,t):
        p=s.d.add_paragraph(); p.paragraph_format.space_before=Pt(9); p.paragraph_format.space_after=Pt(4)
        s._c(p.add_run(f"{n}. {t}"),13.5,True,(0x11,0x33,0x55))
        pr=p._p.get_or_add_pPr(); pb=OxmlElement("w:pBdr"); bt=OxmlElement("w:bottom")
        for k,v in(("w:val","single"),("w:sz","12"),("w:space","2"),("w:color","2B6CB0")): bt.set(qn(k),v)
        pb.append(bt); pr.append(pb)
    def title(s,t,sub): s.para(t,20,True,(0x11,0x33,0x55),2); s.para(sub,12,False,(0x55,0x55,0x55),8)
    def table(s,heads,rows,ctr=set(),fs=8):
        t=s.d.add_table(rows=1,cols=len(heads)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for i,h in enumerate(heads):
            c=t.rows[0].cells[i]; c.text=""; s._c(c.paragraphs[0].add_run(h),fs,True,(0xff,0xff,0xff))
            c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
            sh=OxmlElement("w:shd"); sh.set(qn("w:fill"),"1F2937"); c._tc.get_or_add_tcPr().append(sh)
        for row in rows:
            cs=t.add_row().cells
            for i,v in enumerate(row):
                cs[i].text=""; s._c(cs[i].paragraphs[0].add_run(str(v)),fs)
                cs[i].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER if i in ctr else WD_ALIGN_PARAGRAPH.LEFT
        return t
    def save(s,nm): p=f"{OUTD}/{nm}.docx"; s.d.save(p); print("DOCX:",p)

MSPEC={"PEEK":("PEEK (CF-PEEK 권장)·어닐링","PEEK(建议CF-PEEK)·退火","390~410°C / 130°C / 챔버 60~90°C / 인필 60~100% / 0.15~0.2mm"),
       "PETG":("PETG","PETG","235~245°C / 80°C / 인필 30~40% / 0.2mm"),
       "RIG":("PETG 또는 PLA+ (카메라 거치대)","PETG或PLA+(相机支架)","240/210°C / 80/60°C / 인필 25~35% / 0.24~0.3mm")}
def build(lang):
    D=Doc(lang); T=D.T
    D.title(T("3D 프린트 부품 명세 (최종) + STL 패키지","3D打印件清单(最终)+STL包"),
            T("전 시스템 3D프린트 부품 · 개별분리·중복통합·홀·나사 (LBH-3DP-001 Rev.B)",
              "全系统3D打印件·分离·去重·孔·螺纹 (LBH-3DP-001 Rev.B)"))
    D.table([T("항목","项目"),T("내용","内容")],
      [[T("문서번호","文件号"),"LBH-3DP-001 Rev.B"],[T("일자","日期"),"2026-07-14"],
       [T("STL 패키지","STL包"),f"LBH-3DP-PARTS-FINAL.zip ({len(MAN)} STL, "+T("총","共")+f" {tot_q} ea, "+T("전량 수밀","全部水密")+")"],
       [T("좌표·단위","坐标·单位"),T("시뮬 월드좌표 실측 지오메트리, mm","仿真世界坐标实测几何,mm")],
       [T("재질 분류","材质分类"),T("hint 표기(PEEK/PETG/RIG) — 최종 분류는 발주측","仅标注(PEEK/PETG/RIG)—最终由甲方分类")]],ctr={0})
    D.h1(1,T("재료별 요약","按材料汇总"))
    D.table([T("재료","材料"),T("고유부품","唯一件"),T("총수량","总数量")],
      [[m,bymat[m]["n"],bymat[m]["q"]] for m in bymat],ctr={0,1,2})
    D.h1(2,T("재료별 프린트 사양","按材料打印规格"))
    D.table([T("재료","材料"),T("사양","规格")],
      [[T(MSPEC[m][0],MSPEC[m][1]),MSPEC[m][2]] for m in MSPEC if m in bymat],ctr={0})
    D.h1(3,T("적용 규격 (통일)","应用规格(统一)"))
    D.para(T("· 섬스크류: 1종 통일 — 헤드 Ø12×6(널링 18각) + 프린트용 굵은 나사 M5×P2.0(금속 0.8 대비 2.5배), 나사부 전체 구현.\n"
             "· 섬스크류 소켓/너트: 동일 P2.0 굵은 암나사(체결여유 +0.4).\n"
             "· 스크류 홀: 실계측 위치·기능별 — 관통(THROUGH)=clearHole(M5 Ø5.5/M6 Ø6.6/M10 Ø11) / 탭(TAP)=tapDrill(M5 Ø4.2/M6 Ø5.0) 블라인드.\n"
             "· P01(버킷함+반깔대기)·P15(튜브)는 코드 실측치수로 파라메트릭 재구성(솔리드).",
             "· 拇指螺钉: 统一1种 — 头Ø12×6(滚花18角)+打印粗牙M5×P2.0(金属0.8的2.5倍),全螺纹实现。\n"
             "· 螺钉座/螺母: 同P2.0粗内牙(配合余量+0.4)。\n"
             "· 螺钉孔: 实测位置·按功能 — 通孔(THROUGH)=clearHole / 攻牙(TAP)=tapDrill 盲孔。\n"
             "· P01(料斗盒+半漏斗)·P15(管)按代码实测尺寸参数化重建(实体)。"),9)
    D.d.add_page_break()
    D.h1(4,T("부품 목록 (전량)","部件目录(全量)"))
    rows=[]
    for x in MAN:
        rows.append([x["part"],x.get("dims","-"),x["material_hint"],x["qty"],
                     x.get("holes","-"),x.get("thumb_sockets","-"),T(x.get("status","-"),x.get("status","-"))])
    D.table([T("부품","件"),T("외형(mm)","外形"),T("재료hint","材料"),T("수량","数量"),
             T("홀","孔"),T("소켓","座"),T("상태","状态")],rows,ctr={0,3,4,5},fs=7.5)
    D.para(T("※ 파일명에 수량·외형·홀·소켓 표기(예: P25_qty8_..._h0_ts0). 동일 부품은 1 STL + 수량.",
             "※ 文件名含数量·外形·孔·座(例 P25_qty8...)。相同件1 STL+数量。"),8,False,(0x55,0x55,0x55))
    D.save(f"LBH-3DP-001_{'KR' if lang=='kr' else 'CN'}")
build("kr"); build("cn")
print("done. parts",len(MAN),"total qty",tot_q,"bymat",{k:v for k,v in bymat.items()})
